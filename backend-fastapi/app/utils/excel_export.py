from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RELLENO_ENCABEZADO = PatternFill(start_color="0D0D0D", end_color="0D0D0D", fill_type="solid")
FUENTE_ENCABEZADO = Font(color="FFFFFF", bold=True)
FUENTE_NARANJA = Font(color="FF6600", bold=True, size=13)


def generar_excel_reporte_diario(fecha: str, filas: list[dict], totales: dict) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte diario"

    hoja.merge_cells("A1:G1")
    hoja["A1"] = f"KTM Catálogo — Reporte diario de solicitudes/ventas ({fecha})"
    hoja["A1"].font = FUENTE_NARANJA

    encabezados = ["N° Venta", "Fecha/Hora", "Cliente", "Productos", "Cantidad", "Total", "Estado"]
    hoja.append([])
    hoja.append(encabezados)
    for columna in range(1, len(encabezados) + 1):
        celda = hoja.cell(row=3, column=columna)
        celda.fill = RELLENO_ENCABEZADO
        celda.font = FUENTE_ENCABEZADO
        celda.alignment = Alignment(horizontal="center")

    for fila in filas:
        hoja.append(
            [
                fila["numero_venta"],
                fila["fecha_hora"],
                fila["cliente"],
                fila["productos"],
                fila["cantidad_total"],
                float(fila["total"]),
                fila["estado"],
            ]
        )

    fila_totales = hoja.max_row + 2
    hoja.cell(row=fila_totales, column=1, value="Totales").font = Font(bold=True)
    hoja.cell(row=fila_totales, column=5, value=totales["total_unidades"]).font = Font(bold=True)
    hoja.cell(row=fila_totales, column=6, value=float(totales["total_recaudado"])).font = Font(bold=True)
    hoja.cell(row=fila_totales + 1, column=1, value=f"Total de solicitudes: {totales['total_solicitudes']}")

    for columna in range(1, len(encabezados) + 1):
        hoja.column_dimensions[get_column_letter(columna)].width = 22

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
